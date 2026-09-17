"""
ClassFellow Web - Financial Audit Excel Export Engine
Generates styled multi-column XLSX spreadsheets for institutional monthly fee collections.
"""

from decimal import Decimal
import io
from typing import Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.core.models import AcademicSession, Campus, InstitutionProfile
from apps.fees.models import FeeInvoice, PaymentStatus
from apps.fees.services import FeeWebService


class FeeExcelExportService:
    """Encapsulates styled openpyxl Excel spreadsheet exports for financial audits."""

    @staticmethod
    def export_monthly_collection_workbook(
        session_id: int,
        month_year: str,
        campus_id: Optional[int] = None,
    ) -> io.BytesIO:
        """
        Generates a styled multi-column Excel spreadsheet for fee collection audit.
        Columns:
          Admission #, Student Name, Father Name, Class & Section, Tuition Fee,
          Total Billed, Discount, Total Paid, Balance, Status, Last Receipt #
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Collection {month_year}"

        # Ensure grid lines visible
        ws.views.sheetView[0].showGridLines = True

        # Institutional Metadata Header
        institution = InstitutionProfile.objects.first()
        inst_name = institution.name if institution else "ClassFellow High School & Academy"
        session = AcademicSession.objects.filter(id=session_id).first()
        session_name = session.name if session else f"Session #{session_id}"

        campus = Campus.objects.filter(id=campus_id).first() if campus_id else None
        campus_title = f" | Campus: {campus.name}" if campus else " | All Campuses"

        ws.merge_cells("A1:K1")
        title_cell = ws["A1"]
        title_cell.value = f"{inst_name.upper()} — MONTHLY FEE COLLECTION AUDIT"
        title_cell.font = Font(name="Calibri", size=14, bold=True, color="1E293B")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 26

        ws.merge_cells("A2:K2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Billing Cycle: {month_year} | Academic Session: {session_name}{campus_title}"
        sub_cell.font = Font(name="Calibri", size=10, italic=True, color="64748B")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 18

        ws.append([])  # Blank spacer row 3

        # Table Column Headers (Row 4)
        headers = [
            "Admission #",
            "Student Name",
            "Father Name",
            "Class & Section",
            "Tuition Fee",
            "Total Billed",
            "Discount",
            "Total Paid",
            "Balance",
            "Status",
            "Last Receipt #",
        ]
        ws.append(headers)
        ws.row_dimensions[4].height = 24

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(fill_type="solid", start_color="1E293B", end_color="1E293B")
        header_border = Border(
            left=Side(style="thin", color="334155"),
            right=Side(style="thin", color="334155"),
            top=Side(style="medium", color="0F172A"),
            bottom=Side(style="medium", color="0F172A"),
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = header_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Query fee invoices
        invoices_qs = (
            FeeInvoice.objects.filter(
                session_id=session_id,
                month_year=month_year,
            )
            .select_related(
                "enrollment__student",
                "enrollment__class_group",
                "enrollment__class_group__campus",
            )
            .prefetch_related("items", "payments")
            .order_by("enrollment__class_group__name", "enrollment__roll_number", "enrollment__student__first_name")
        )

        if campus_id is not None:
            invoices_qs = invoices_qs.filter(
                enrollment__class_group__campus_id=campus_id
            )

        row_start = 5
        current_row = row_start

        thin_side = Side(style="thin", color="CBD5E1")
        row_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        alt_fill = PatternFill(fill_type="solid", start_color="F8FAFC", end_color="F8FAFC")

        total_tuition_sum = Decimal("0.00")
        total_billed_sum = Decimal("0.00")
        total_discount_sum = Decimal("0.00")
        total_paid_sum = Decimal("0.00")
        total_balance_sum = Decimal("0.00")

        for inv in invoices_qs:
            student = inv.enrollment.student
            cg = inv.enrollment.class_group
            bal = FeeWebService.calculate_invoice_balance(inv)

            tuition_item = next((item for item in inv.items.all() if "tuition" in item.fee_head.name.lower()), None)
            tuition_amt = tuition_item.amount if tuition_item else inv.total_payable

            last_payment = inv.payments.filter(status=PaymentStatus.ISSUED).order_by("-id").first()
            last_receipt = last_payment.receipt_number if last_payment else "—"

            ws.append([
                student.admission_number,
                f"{student.first_name} {student.last_name}".strip(),
                student.guardian_name,
                f"{cg.name} - {cg.section_or_batch}",
                float(tuition_amt),
                float(inv.net_due),
                float(inv.discount_amount),
                float(bal["total_paid"]),
                float(bal["current_balance"]),
                bal["status"],
                last_receipt,
            ])

            total_tuition_sum += tuition_amt
            total_billed_sum += inv.net_due
            total_discount_sum += inv.discount_amount
            total_paid_sum += bal["total_paid"]
            total_balance_sum += bal["current_balance"]

            is_even = (current_row % 2 == 0)
            ws.row_dimensions[current_row].height = 20

            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=current_row, column=col_idx)
                c.font = Font(name="Calibri", size=10)
                c.border = row_border
                if is_even:
                    c.fill = alt_fill

                # Alignment & Number formatting
                if col_idx in (5, 6, 7, 8, 9):
                    c.number_format = "#,##0.00"
                    c.alignment = Alignment(horizontal="right", vertical="center")
                elif col_idx in (1, 10, 11):
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

            current_row += 1

        # Summary Row
        summary_row = current_row
        ws.append([
            "TOTAL",
            f"{invoices_qs.count()} Invoices",
            "",
            "",
            float(total_tuition_sum),
            float(total_billed_sum),
            float(total_discount_sum),
            float(total_paid_sum),
            float(total_balance_sum),
            "",
            "",
        ])
        ws.row_dimensions[summary_row].height = 24

        summary_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="1E293B"),
            bottom=Side(style="double", color="1E293B"),
        )
        summary_fill = PatternFill(fill_type="solid", start_color="E2E8F0", end_color="E2E8F0")

        for col_idx in range(1, len(headers) + 1):
            sc = ws.cell(row=summary_row, column=col_idx)
            sc.font = Font(name="Calibri", size=10, bold=True)
            sc.fill = summary_fill
            sc.border = summary_border
            if col_idx in (5, 6, 7, 8, 9):
                sc.number_format = "#,##0.00"
                sc.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in (1, 2):
                sc.alignment = Alignment(horizontal="left", vertical="center")
            else:
                sc.alignment = Alignment(horizontal="center", vertical="center")

        # Column Auto-Width
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Ignore merged banner row 1 & 2 for width calculation
                if cell.row in (1, 2, 3):
                    continue
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output_buffer = io.BytesIO()
        wb.save(output_buffer)
        output_buffer.seek(0)
        return output_buffer
