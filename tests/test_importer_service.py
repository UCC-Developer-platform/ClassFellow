"""
ClassFellow - Student & Enrollment Bulk Importer Test Suite (tests/test_importer_service.py)
===========================================================================================
Validates Excel (.xlsx) template generation, multi-row Excel & CSV ingestion,
Pakistani phone normalization, admission number validation/generation, discount parsing,
validation error handling, and database integrity.
"""

import os
import csv
import pytest
import sqlite3
from decimal import Decimal
import openpyxl

from services.schema_service import migrate_to_latest
from services.student_service import StudentService
from services.importer_service import StudentImporterService


@pytest.fixture
def migrated_db(db_connection):
    """Provides a SQLite connection with all migrations applied."""
    migrate_to_latest(db_connection)
    return db_connection


@pytest.fixture
def test_context(migrated_db):
    """Sets up an academic session and class group for testing."""
    student_service = StudentService(migrated_db)
    session_id = student_service.create_academic_session(
        name="2026-2027",
        start_date="2026-04-01",
        end_date="2027-03-31"
    )
    class_group_id = student_service.create_class_group(
        session_id=session_id,
        name="Grade 10",
        section_or_batch="Section Rose",
        group_type="SchoolClass",
        monthly_tuition_fee=Decimal("5000.00")
    )
    importer = StudentImporterService(migrated_db)
    return {
        "conn": migrated_db,
        "session_id": session_id,
        "class_group_id": class_group_id,
        "importer": importer,
        "student_service": student_service,
    }


# =============================================================================
# 1. Template Generation Tests
# =============================================================================

def test_generate_excel_template(tmp_path, test_context):
    """Verifies that generate_excel_template produces a valid .xlsx file with expected sheets and headers."""
    importer = test_context["importer"]
    template_path = str(tmp_path / "ClassFellow_Student_Template.xlsx")

    res_path = importer.generate_excel_template(template_path)
    assert os.path.exists(res_path)
    assert res_path == os.path.abspath(template_path)

    wb = openpyxl.load_workbook(res_path)
    assert "Student Roster" in wb.sheetnames
    assert "Instructions" in wb.sheetnames

    # Check Sheet 1 headers
    ws_roster = wb["Student Roster"]
    headers = [cell.value for cell in ws_roster[1]]
    assert headers == StudentImporterService.TEMPLATE_COLUMNS

    # Check sample rows exist
    assert ws_roster.max_row >= 4  # Header + 3 sample rows
    first_sample_first_name = ws_roster.cell(row=2, column=2).value
    assert first_sample_first_name == "Usman"

    # Check Sheet 2 headers
    ws_info = wb["Instructions"]
    info_headers = [cell.value for cell in ws_info[1]]
    assert info_headers == ["Field Name", "Required", "Format & Validation Rules"]
    wb.close()


# =============================================================================
# 2. Multi-Row Excel Import & Database Integrity
# =============================================================================

def test_import_students_from_excel_success(tmp_path, test_context):
    """Tests clean bulk ingestion of students from an Excel workbook with Urdu text and phones."""
    importer = test_context["importer"]
    conn = test_context["conn"]
    session_id = test_context["session_id"]
    class_group_id = test_context["class_group_id"]

    # Create test excel file
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students"

    headers = [
        "Roll Number", "First Name *", "Last Name", "Urdu Name",
        "Gender *", "Guardian Name *", "Guardian Phone *",
        "Admission Number", "Date of Birth", "Address", "Monthly Discount"
    ]
    ws.append(headers)

    rows = [
        ["01", "Zaid", "Khan", "زید خان", "Male", "Tariq Khan", "0300-1122334", "", "2012-05-14", "Gulberg, Lahore", "0.00"],
        ["02", "Ayesha", "Bibi", "عائشہ بی بی", "Female", "Babar Bibi", "+92 321 9988776", "CF-2026-9001", "2012-08-20", "Model Town, Lahore", "500.00"],
        ["03", "Bilal", "", "بلال", "Male", "Sajid Mehmood", "0333 4455667", "", "2011-11-03", "", "200.50"],
    ]
    for r in rows:
        ws.append(r)

    excel_file = str(tmp_path / "students_import.xlsx")
    wb.save(excel_file)
    wb.close()

    result = importer.import_students_from_excel(
        file_path=excel_file,
        class_group_id=class_group_id,
        session_id=session_id
    )

    assert result["total_rows"] == 3
    assert result["imported_count"] == 3
    assert result["failed_count"] == 0
    assert len(result["errors"]) == 0

    # Verify database state in `students`
    cursor = conn.cursor()
    cursor.execute("SELECT first_name, last_name, urdu_name, gender, guardian_phone, admission_number FROM students ORDER BY id ASC;")
    db_students = cursor.fetchall()
    assert len(db_students) == 3

    # Row 1: Zaid Khan, normalized phone
    assert db_students[0][0] == "Zaid"
    assert db_students[0][1] == "Khan"
    assert db_students[0][2] == "زید خان"
    assert db_students[0][3] == "Male"
    assert db_students[0][4] == "03001122334"
    assert db_students[0][5].startswith("CF-")

    # Row 2: Explicit admission CF-2026-9001
    assert db_students[1][0] == "Ayesha"
    assert db_students[1][4] == "03219988776"
    assert db_students[1][5] == "CF-2026-9001"

    # Row 3: Bilal
    assert db_students[2][0] == "Bilal"
    assert db_students[2][1] is None
    assert db_students[2][4] == "03334455667"

    # Verify enrollments table
    cursor.execute("SELECT student_id, class_group_id, session_id, roll_number, custom_discount_amount FROM enrollments ORDER BY id ASC;")
    db_enrollments = cursor.fetchall()
    assert len(db_enrollments) == 3
    assert db_enrollments[0][1] == class_group_id
    assert db_enrollments[0][2] == session_id
    assert db_enrollments[0][3] == "01"
    assert Decimal(db_enrollments[0][4]) == Decimal("0.00")

    assert db_enrollments[1][3] == "02"
    assert Decimal(db_enrollments[1][4]) == Decimal("500.00")

    assert db_enrollments[2][3] == "03"
    assert Decimal(db_enrollments[2][4]) == Decimal("200.50")


# =============================================================================
# 3. CSV File Import Tests
# =============================================================================

def test_import_students_from_csv(tmp_path, test_context):
    """Tests student ingestion from a UTF-8 CSV spreadsheet."""
    importer = test_context["importer"]
    conn = test_context["conn"]
    session_id = test_context["session_id"]
    class_group_id = test_context["class_group_id"]

    csv_file = str(tmp_path / "students.csv")
    with open(csv_file, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Roll", "First Name", "Last Name", "Urdu Name", "Gender", "Guardian", "Mobile", "Admission No", "DOB", "Address", "Discount"])
        writer.writerow(["10", "Hamza", "Abbasi", "حمزہ عباسی", "Male", "Farooq Abbasi", "0312-3344556", "", "2013-02-10", "Islamabad", "0"])
        writer.writerow(["11", "Sadia", "Akram", "سعدیہ اکرم", "Female", "Akram Sheikh", "+923456789012", "", "2013-06-15", "Rawalpindi", "100.00"])

    result = importer.import_students_from_excel(
        file_path=csv_file,
        class_group_id=class_group_id,
        session_id=session_id
    )

    assert result["total_rows"] == 2
    assert result["imported_count"] == 2
    assert result["failed_count"] == 0

    cursor = conn.cursor()
    cursor.execute("SELECT first_name, guardian_phone FROM students WHERE first_name IN ('Hamza', 'Sadia');")
    rows = dict(cursor.fetchall())
    assert rows["Hamza"] == "03123344556"
    assert rows["Sadia"] == "03456789012"


# =============================================================================
# 4. Validation Errors & Partial Failure Handling
# =============================================================================

def test_import_validation_errors(tmp_path, test_context):
    """Verifies that invalid rows are reported with row numbers while valid rows succeed."""
    importer = test_context["importer"]
    session_id = test_context["session_id"]
    class_group_id = test_context["class_group_id"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["First Name *", "Gender *", "Guardian Name *", "Guardian Phone *", "Admission Number", "Monthly Discount"])

    # Row 2: Valid
    ws.append(["Kamran", "Male", "Naveed", "03001234567", "", "0.00"])
    # Row 3: Missing first name
    ws.append(["", "Male", "Rashid", "03001234568", "", "0.00"])
    # Row 4: Invalid phone number
    ws.append(["Zain", "Male", "Arshad", "04235889900", "", "0.00"])
    # Row 5: Invalid gender
    ws.append(["Maria", "UnknownGender", "Sohail", "03001234569", "", "0.00"])
    # Row 6: Negative discount
    ws.append(["Omar", "Male", "Javed", "03001234570", "", "-100.00"])
    # Row 7: Valid
    ws.append(["Sara", "Female", "Altaf", "03001234571", "", "150.00"])

    excel_file = str(tmp_path / "errors.xlsx")
    wb.save(excel_file)
    wb.close()

    result = importer.import_students_from_excel(excel_file, class_group_id, session_id)

    assert result["total_rows"] == 6
    assert result["imported_count"] == 2  # Kamran & Sara
    assert result["failed_count"] == 4
    assert len(result["errors"]) == 4

    # Check error details
    err_rows = {e["row"]: e["error"] for e in result["errors"]}
    assert 3 in err_rows
    assert "first name is mandatory" in err_rows[3].lower()

    assert 4 in err_rows
    assert "invalid phone format" in err_rows[4].lower()

    assert 5 in err_rows
    assert "invalid gender" in err_rows[5].lower()

    assert 6 in err_rows
    assert "discount cannot be negative" in err_rows[6].lower()


def test_import_duplicate_admission_numbers(tmp_path, test_context):
    """Verifies duplicate admission numbers within the spreadsheet or against DB are rejected."""
    importer = test_context["importer"]
    session_id = test_context["session_id"]
    class_group_id = test_context["class_group_id"]

    # First insert a student with admission number CF-2026-5555
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.append(["First Name *", "Gender *", "Guardian Name *", "Guardian Phone *", "Admission Number"])
    ws1.append(["ExistingStudent", "Male", "Guardian A", "03001111111", "CF-2026-5555"])
    f1 = str(tmp_path / "sheet1.xlsx")
    wb1.save(f1)
    wb1.close()

    res1 = importer.import_students_from_excel(f1, class_group_id, session_id)
    assert res1["imported_count"] == 1

    # Second sheet: row with existing CF-2026-5555, and two rows repeating CF-2026-7777
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.append(["First Name *", "Gender *", "Guardian Name *", "Guardian Phone *", "Admission Number"])
    ws2.append(["Student1", "Male", "Guardian B", "03002222222", "CF-2026-5555"])  # Duplicate with DB
    ws2.append(["Student2", "Female", "Guardian C", "03003333333", "CF-2026-7777"]) # First seen
    ws2.append(["Student3", "Female", "Guardian D", "03004444444", "CF-2026-7777"]) # Duplicate in sheet

    f2 = str(tmp_path / "sheet2.xlsx")
    wb2.save(f2)
    wb2.close()

    res2 = importer.import_students_from_excel(f2, class_group_id, session_id)
    assert res2["total_rows"] == 3
    assert res2["imported_count"] == 1  # Student2 succeeds
    assert res2["failed_count"] == 2

    err_msgs = [e["error"].lower() for e in res2["errors"]]
    assert any("already registered in database" in m for m in err_msgs)
    assert any("repeated within spreadsheet" in m for m in err_msgs)


# =============================================================================
# 5. Boundary Conditions & Foreign Keys
# =============================================================================

def test_import_invalid_foreign_keys_and_files(tmp_path, test_context):
    """Verifies proper error handling when foreign keys or files are invalid."""
    importer = test_context["importer"]
    session_id = test_context["session_id"]
    class_group_id = test_context["class_group_id"]

    # Non-existent file
    with pytest.raises(FileNotFoundError):
        importer.import_students_from_excel("non_existent_file.xlsx", class_group_id, session_id)

    # Invalid session
    valid_file = str(tmp_path / "dummy.xlsx")
    wb = openpyxl.Workbook()
    wb.save(valid_file)
    wb.close()

    with pytest.raises(ValueError, match="Academic session with id=999999 does not exist"):
        importer.import_students_from_excel(valid_file, class_group_id, 999999)

    # Invalid class group
    with pytest.raises(ValueError, match="Class group id=999999 does not exist"):
        importer.import_students_from_excel(valid_file, 999999, session_id)

    # Unsupported file extension
    bad_ext_file = str(tmp_path / "data.txt")
    with open(bad_ext_file, "w") as f:
        f.write("some text")
    with pytest.raises(ValueError, match="Unsupported file format"):
        importer.import_students_from_excel(bad_ext_file, class_group_id, session_id)


def test_import_empty_file(tmp_path, test_context):
    """Verifies that an empty sheet returns 0 counts with no errors."""
    importer = test_context["importer"]
    session_id = test_context["session_id"]
    class_group_id = test_context["class_group_id"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(StudentImporterService.TEMPLATE_COLUMNS)
    empty_excel = str(tmp_path / "empty.xlsx")
    wb.save(empty_excel)
    wb.close()

    result = importer.import_students_from_excel(empty_excel, class_group_id, session_id)
    assert result["total_rows"] == 0
    assert result["imported_count"] == 0
    assert result["failed_count"] == 0
    assert result["errors"] == []
